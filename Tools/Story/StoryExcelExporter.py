from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# 경로
# ============================================================

SCRIPT_PATH = Path(__file__).resolve()
PROJECT_ROOT = SCRIPT_PATH.parents[2]

STORY_DATA_FOLDER = (
    PROJECT_ROOT
    / "Assets"
    / "PHA"
    / "Story"
    / "Data"
)

EXCEL_FOLDER = STORY_DATA_FOLDER / "Excel"
CSV_FOLDER = STORY_DATA_FOLDER / "CSV"

EXCEL_PATH = EXCEL_FOLDER / "StoryData.xlsx"


STORIES_CSV = CSV_FOLDER / "Stories.csv"
NODES_CSV = CSV_FOLDER / "StoryNodes.csv"
CHOICES_CSV = CSV_FOLDER / "StoryChoices.csv"
CONDITIONS_CSV = CSV_FOLDER / "StoryConditions.csv"
EFFECTS_CSV = CSV_FOLDER / "StoryEffects.csv"
CHARACTERS_CSV = CSV_FOLDER / "Characters.csv"


# ============================================================
# 현재 Excel 시트 구조
# ============================================================

META_HEADER_ROW = 4
META_DATA_ROW = 5

NODE_HEADER_ROW = 11
NODE_START_ROW = 12
NODE_END_ROW = 220

CHOICE_HEADER_ROW = 225
CHOICE_START_ROW = 226
CHOICE_END_ROW = 300

CONDITION_HEADER_ROW = 305
CONDITION_START_ROW = 306
CONDITION_END_ROW = 340

EFFECT_HEADER_ROW = 345
EFFECT_START_ROW = 346
EFFECT_END_ROW = 400


# ============================================================
# 제외할 시스템 시트
# ============================================================

SYSTEM_SHEETS = {
    "README",
    "Enums",
    "Characters",
    "StoryList",
    "STORY_TEMPLATE",
}


# ============================================================
# CSV Header
# ============================================================

STORY_HEADERS = [
    "StoryId",
    "StoryType",
    "StoryTitle",
    "Year",
    "Month",
    "Day",
    "Priority",
    "PlayOnce",
    "PlayTiming",
    "UnlockFactionId",
    "StartNodeId",
    "Memo",
]

NODE_HEADERS = [
    "StoryId",
    "NodeOrder",
    "NodeId",
    "NodeType",
    "Text",
    "CharacterId",
    "PortraitId",
    "NextNodeId",
    "DimPortrait",
    "KeepPortrait",
    "UseTypingEffect",
    "AutoAdvance",
    "AutoAdvanceDelay",
    "Memo",
]

CHOICE_HEADERS = [
    "StoryId",
    "NodeId",
    "ChoiceIndex",
    "ChoiceText",
    "TargetNodeId",
    "ResultKey",
    "ResultValue",
    "UseCondition",
    "RequiredKey",
    "RequiredValue",
    "HideWhenLocked",
    "LockedText",
    "Memo",
]

CONDITION_HEADERS = [
    "StoryId",
    "ConditionIndex",
    "ConditionType",
    "Key",
    "Value",
    "IntValue",
    "Memo",
]

EFFECT_HEADERS = [
    "StoryId",
    "NodeId",
    "EffectIndex",
    "EffectType",
    "Target",
    "Duration",
    "Strength",
    "WaitForCompletion",
    "Memo",
]

CHARACTER_HEADERS = [
    "CharacterId",
    "CharacterName",
    "CharacterNameEng",
    "PortraitId",
    "PortraitLabel",
    "SpriteName",
]


# ============================================================
# 공통 함수
# ============================================================

def cell_value(
    worksheet,
    row: int,
    column: int,
    default=""
):
    value = worksheet.cell(
        row=row,
        column=column
    ).value

    if value is None:
        return default

    return value


def text_value(value):
    if value is None:
        return ""

    return str(value)


def normalize_bool(value):
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    text = str(value).strip().upper()

    if text in (
        "TRUE",
        "1",
        "YES",
        "Y",
    ):
        return "TRUE"

    return "FALSE"


def is_empty_row(values):
    return not any(
        value is not None
        and str(value).strip() != ""
        for value in values
    )


def write_csv(
    path: Path,
    headers: list[str],
    rows: list[list]
):
    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with path.open(
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        writer = csv.writer(
            file,
            quoting=csv.QUOTE_MINIMAL
        )

        writer.writerow(headers)

        writer.writerows(rows)

    print(
        f"[CSV] {path.name} "
        f"({len(rows)} rows)"
    )


# ============================================================
# Story 기본 정보
# ============================================================

def read_story_meta(
    worksheet
):
    values = [
        cell_value(
            worksheet,
            META_DATA_ROW,
            column
        )
        for column in range(
            1,
            13
        )
    ]

    story_id = text_value(
        values[0]
    ).strip()

    if not story_id:
        return None

    return [
        story_id,

        text_value(
            values[1]
        ).strip(),

        text_value(
            values[2]
        ),

        values[3] or 0,
        values[4] or 0,
        values[5] or 0,

        values[6] or 100,

        normalize_bool(
            values[7]
        ),

        text_value(
            values[8]
        ).strip(),

        text_value(
            values[9]
        ).strip(),

        text_value(
            values[10]
        ).strip(),

        text_value(
            values[11]
        ),
    ]


# ============================================================
# Nodes
# ============================================================

def read_nodes(
    worksheet,
    story_id
):
    rows = []

    for row in range(
        NODE_START_ROW,
        NODE_END_ROW + 1
    ):
        node_id = text_value(
            cell_value(
                worksheet,
                row,
                2
            )
        ).strip()

        if not node_id:
            continue

        node_type = text_value(
            cell_value(
                worksheet,
                row,
                3
            )
        ).strip()

        rows.append(
            [
                story_id,

                cell_value(
                    worksheet,
                    row,
                    1,
                    0
                ),

                node_id,

                node_type,

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        4
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        5
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        6
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        7
                    )
                ).strip(),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        8
                    )
                ),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        9
                    )
                ),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        10
                    )
                ),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        11
                    )
                ),

                cell_value(
                    worksheet,
                    row,
                    12,
                    1.5
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        14
                    )
                ),
            ]
        )

    return rows


# ============================================================
# Choices
# ============================================================

def read_choices(
    worksheet,
    story_id
):
    rows = []

    for row in range(
        CHOICE_START_ROW,
        CHOICE_END_ROW + 1
    ):
        choice_node_id = text_value(
            cell_value(
                worksheet,
                row,
                1
            )
        ).strip()

        if not choice_node_id:
            continue

        rows.append(
            [
                story_id,
                choice_node_id,

                cell_value(
                    worksheet,
                    row,
                    2,
                    0
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        3
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        4
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        5
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        6
                    )
                ).strip(),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        7
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        8
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        9
                    )
                ).strip(),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        10
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        11
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        12
                    )
                ),
            ]
        )

    return rows


# ============================================================
# Conditions
# ============================================================

def read_conditions(
    worksheet,
    story_id
):
    rows = []

    for row in range(
        CONDITION_START_ROW,
        CONDITION_END_ROW + 1
    ):
        condition_type = text_value(
            cell_value(
                worksheet,
                row,
                2
            )
        ).strip()

        if not condition_type:
            continue

        # None은 실제 조건 없음으로 처리
        if condition_type == "None":
            continue

        rows.append(
            [
                story_id,

                cell_value(
                    worksheet,
                    row,
                    1,
                    0
                ),

                condition_type,

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        3
                    )
                ).strip(),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        4
                    )
                ).strip(),

                cell_value(
                    worksheet,
                    row,
                    5,
                    0
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        6
                    )
                ),
            ]
        )

    return rows


# ============================================================
# Effects
# ============================================================

def read_effects(
    worksheet,
    story_id
):
    rows = []

    for row in range(
        EFFECT_START_ROW,
        EFFECT_END_ROW + 1
    ):
        node_id = text_value(
            cell_value(
                worksheet,
                row,
                1
            )
        ).strip()

        effect_type = text_value(
            cell_value(
                worksheet,
                row,
                3
            )
        ).strip()

        if not node_id:
            continue

        if (
            not effect_type
            or effect_type == "None"
        ):
            continue

        rows.append(
            [
                story_id,
                node_id,

                cell_value(
                    worksheet,
                    row,
                    2,
                    0
                ),

                effect_type,

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        4
                    )
                ).strip(),

                cell_value(
                    worksheet,
                    row,
                    5,
                    0
                ),

                cell_value(
                    worksheet,
                    row,
                    6,
                    1
                ),

                normalize_bool(
                    cell_value(
                        worksheet,
                        row,
                        7
                    )
                ),

                text_value(
                    cell_value(
                        worksheet,
                        row,
                        8
                    )
                ),
            ]
        )

    return rows


# ============================================================
# Characters
# ============================================================

def read_characters(
    workbook
):
    if "Characters" not in workbook.sheetnames:
        return []

    worksheet = workbook[
        "Characters"
    ]

    rows = []

    row = 2

    while True:
        values = [
            cell_value(
                worksheet,
                row,
                column
            )
            for column in range(
                1,
                7
            )
        ]

        if is_empty_row(values):
            break

        character_id = text_value(
            values[0]
        ).strip()

        if character_id:
            rows.append(
                [
                    character_id,
                    text_value(values[1]),
                    text_value(values[2]),
                    text_value(values[3]),
                    text_value(values[4]),
                    text_value(values[5]),
                ]
            )

        row += 1

    return rows


# ============================================================
# Excel → CSV
# ============================================================

def export_excel():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Excel 파일이 없습니다:\n"
            f"{EXCEL_PATH}"
        )

    print("=" * 60)
    print("Story Excel Exporter")
    print("=" * 60)

    workbook = load_workbook(
        EXCEL_PATH,
        data_only=False
    )

    stories = []
    nodes = []
    choices = []
    conditions = []
    effects = []

    story_ids = set()

    # --------------------------------------------------------
    # 스토리 시트 탐색
    # --------------------------------------------------------

    for sheet_name in workbook.sheetnames:

        if sheet_name in SYSTEM_SHEETS:
            continue

        worksheet = workbook[
            sheet_name
        ]

        meta = read_story_meta(
            worksheet
        )

        if meta is None:
            print(
                f"[SKIP] StoryId 없음: "
                f"{sheet_name}"
            )

            continue

        story_id = meta[0]

        if story_id in story_ids:
            raise ValueError(
                f"중복 StoryId 발견: "
                f"{story_id}"
            )

        story_ids.add(
            story_id
        )

        stories.append(
            meta
        )

        story_nodes = read_nodes(
            worksheet,
            story_id
        )

        story_choices = read_choices(
            worksheet,
            story_id
        )

        story_conditions = (
            read_conditions(
                worksheet,
                story_id
            )
        )

        story_effects = read_effects(
            worksheet,
            story_id
        )

        nodes.extend(
            story_nodes
        )

        choices.extend(
            story_choices
        )

        conditions.extend(
            story_conditions
        )

        effects.extend(
            story_effects
        )

        print(
            f"[Story] {story_id}"
            f" | Nodes {len(story_nodes)}"
            f" | Choices {len(story_choices)}"
        )

    # --------------------------------------------------------
    # Characters
    # --------------------------------------------------------

    characters = read_characters(
        workbook
    )

    # --------------------------------------------------------
    # CSV 저장
    # --------------------------------------------------------

    write_csv(
        STORIES_CSV,
        STORY_HEADERS,
        stories
    )

    write_csv(
        NODES_CSV,
        NODE_HEADERS,
        nodes
    )

    write_csv(
        CHOICES_CSV,
        CHOICE_HEADERS,
        choices
    )

    write_csv(
        CONDITIONS_CSV,
        CONDITION_HEADERS,
        conditions
    )

    write_csv(
        EFFECTS_CSV,
        EFFECT_HEADERS,
        effects
    )

    if characters:
        write_csv(
            CHARACTERS_CSV,
            CHARACTER_HEADERS,
            characters
        )

    print()
    print("=" * 60)
    print("Export 완료")
    print("=" * 60)

    print(
        f"Stories   : {len(stories)}"
    )

    print(
        f"Nodes     : {len(nodes)}"
    )

    print(
        f"Choices   : {len(choices)}"
    )

    print(
        f"Conditions: {len(conditions)}"
    )

    print(
        f"Effects   : {len(effects)}"
    )

    print(
        f"Characters: {len(characters)}"
    )

    print()
    print(
        f"CSV Folder:\n{CSV_FOLDER}"
    )


# ============================================================
# Main
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "StoryData.xlsx를 Unity용 "
            "Story CSV 파일들로 변환합니다."
        )
    )

    parser.parse_args()

    try:
        export_excel()

    except Exception as exception:

        print()
        print(
            "[ERROR] Story Excel Export 실패"
        )

        print(
            str(exception)
        )

        sys.exit(1)


if __name__ == "__main__":
    main()

# 실행 : python StoryExcelExporter.py
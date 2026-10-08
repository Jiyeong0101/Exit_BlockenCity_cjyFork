from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule


# ============================================================
# 경로
# ============================================================

# StoryExcelBuilder.py
# Tools/Story/StoryExcelBuilder.py
#
# parents[0] = Story
# parents[1] = Tools
# parents[2] = 프로젝트 루트

SCRIPT_PATH = Path(__file__).resolve()
PROJECT_ROOT = SCRIPT_PATH.parents[2]

STORY_DATA_FOLDER = (
    PROJECT_ROOT
    / "Assets"
    / "PHA"
    / "Story"
    / "Data"
)

CSV_FOLDER = STORY_DATA_FOLDER / "CSV"
EXCEL_FOLDER = STORY_DATA_FOLDER / "Excel"

EXCEL_PATH = EXCEL_FOLDER / "StoryData.xlsx"


STORIES_CSV = CSV_FOLDER / "Stories.csv"
NODES_CSV = CSV_FOLDER / "StoryNodes.csv"
CHOICES_CSV = CSV_FOLDER / "StoryChoices.csv"
CONDITIONS_CSV = CSV_FOLDER / "StoryConditions.csv"
EFFECTS_CSV = CSV_FOLDER / "StoryEffects.csv"
CHARACTERS_CSV = CSV_FOLDER / "Characters.csv"

# ============================================================
# Excel 디자인
# ============================================================

COLOR_TITLE = "243447"
COLOR_SECTION = "D9E2F3"

COLOR_HEADER = "34495E"
COLOR_HEADER_FONT = "FFFFFF"

COLOR_CHARACTER = "E8F3FF"
COLOR_PLAYER = "E9F7EF"
COLOR_NARRATION = "F4F4F5"
COLOR_CHOICE = "FFF4D6"
COLOR_END = "FDE2E2"

COLOR_WARNING = "F8CBAD"

COLOR_INFO = "EAF2F8"

COLOR_BORDER = "C9CED6"


thin_side = Side(
    style="thin",
    color=COLOR_BORDER
)

thin_border = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side,
)


# ============================================================
# Enum
# ============================================================

STORY_TYPES = [
    "Normal",
    "Force",
    "Unlock",
    "Branch",
    "Event",
]

PLAY_TIMINGS = [
    "BeforeMonthly",
    "Monthly",
    "AfterMonthly",
]

NODE_TYPES = [
    "CharacterDialogue",
    "PlayerDialogue",
    "Narration",
    "Choice",
    "End",
]

CONDITION_TYPES = [
    "None",
    "StoryCompleted",
    "StoryNotCompleted",
    "FactionIntroduced",
    "FactionNotIntroduced",
    "RelationshipAtLeast",
    "RelationshipAtMost",
    "ChoiceEquals",
]

EFFECT_TYPES = [
    "None",
    "PortraitFadeIn",
    "PortraitFadeOut",
    "PortraitShakeSmall",
    "PortraitShakeStrong",
    "PortraitZoomIn",
    "PortraitZoomOut",
    "ScreenFlash",
    "ScreenDarken",
    "TextShake",
]

EFFECT_TARGETS = [
    "Portrait",
    "DialoguePanel",
    "DialogueText",
    "Background",
    "EntireScreen",
]


# ============================================================
# CSV
# ============================================================

def read_csv(path: Path, required: bool = False) -> list[dict]:
    """
    UTF-8 BOM CSV를 읽는다.
    Exporter에서 multiline field를 사용해도 csv 모듈이 처리한다.
    """

    if not path.exists():

        if required:
            raise FileNotFoundError(
                f"필수 CSV가 없습니다: {path}"
            )

        print(
            f"[INFO] 선택 CSV 없음: {path.name}"
        )

        return []

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        reader = csv.DictReader(file)

        result = []

        for row in reader:

            if not row:
                continue

            if not any(
                value and str(value).strip()
                for value in row.values()
            ):
                continue

            cleaned = {}

            for key, value in row.items():

                if key is None:
                    continue

                cleaned[key.strip()] = (
                    value if value is not None else ""
                )

            result.append(cleaned)

        return result


# ============================================================
# 공통 함수
# ============================================================

def get_value(
    row: dict,
    key: str,
    default=""
):
    value = row.get(key, default)

    if value is None:
        return default

    return value


def parse_int(
    value,
    default=0
):
    try:
        return int(str(value).strip())
    except (ValueError, TypeError):
        return default


def parse_float(
    value,
    default=0.0
):
    try:
        return float(str(value).strip())
    except (ValueError, TypeError):
        return default


def normalize_bool(
    value,
    default="FALSE"
):
    if value is None:
        return default

    text = str(value).strip().upper()

    if text in ("TRUE", "1", "YES", "Y"):
        return "TRUE"

    if text in ("FALSE", "0", "NO", "N"):
        return "FALSE"

    return default


def sanitize_sheet_name(name: str) -> str:
    """
    Excel:
    - 최대 31자
    - : \\ / ? * [ ] 사용 불가
    """

    if not name:
        return "Story"

    name = re.sub(
        r'[:\\/?*\[\]]',
        "_",
        name
    )

    name = name.strip()

    if len(name) > 31:
        name = name[:31]

    return name or "Story"


def make_unique_sheet_name(
    workbook: Workbook,
    desired: str
):
    desired = sanitize_sheet_name(desired)

    if desired not in workbook.sheetnames:
        return desired

    base = desired

    index = 2

    while True:

        suffix = f"_{index}"

        max_length = 31 - len(suffix)

        candidate = (
            base[:max_length]
            + suffix
        )

        if candidate not in workbook.sheetnames:
            return candidate

        index += 1


def auto_story_sheet_name(
    story: dict
):
    """
    Main_04 / 4월 / 고요의 순간
        → 04_고요의순간

    Event
        → EV_홍련귀세력소개
    """

    story_id = get_value(
        story,
        "StoryId"
    )

    title = get_value(
        story,
        "StoryTitle",
        story_id
    )

    title_for_sheet = re.sub(
        r"\s+",
        "",
        title
    )

    story_type = get_value(
        story,
        "StoryType"
    )

    month = parse_int(
        get_value(
            story,
            "Month"
        )
    )

    if story_type == "Event":
        return sanitize_sheet_name(
            f"EV_{title_for_sheet}"
        )

    if month > 0:
        return sanitize_sheet_name(
            f"{month:02d}_{title_for_sheet}"
        )

    return sanitize_sheet_name(
        f"{story_id}_{title_for_sheet}"
    )


# ============================================================
# 스타일
# ============================================================

def style_title(
    worksheet,
    cell_range,
    text
):
    worksheet.merge_cells(cell_range)

    start = cell_range.split(":")[0]

    cell = worksheet[start]

    cell.value = text

    cell.fill = PatternFill(
        "solid",
        fgColor=COLOR_TITLE
    )

    cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=15
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[
        cell.row
    ].height = 32


def style_section(
    worksheet,
    row: int,
    text: str
):
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=14
    )

    cell = worksheet.cell(
        row=row,
        column=1
    )

    cell.value = text

    cell.fill = PatternFill(
        "solid",
        fgColor=COLOR_SECTION
    )

    cell.font = Font(
        bold=True,
        color="1F2937",
        size=12
    )

    cell.alignment = Alignment(
        vertical="center"
    )

    worksheet.row_dimensions[
        row
    ].height = 24


def style_headers(
    worksheet,
    row: int,
    headers: list[str]
):
    for column, header in enumerate(
        headers,
        start=1
    ):
        cell = worksheet.cell(
            row=row,
            column=column
        )

        cell.value = header

        cell.fill = PatternFill(
            "solid",
            fgColor=COLOR_HEADER
        )

        cell.font = Font(
            bold=True,
            color=COLOR_HEADER_FONT
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    worksheet.row_dimensions[
        row
    ].height = 30


def apply_standard_cell_style(
    worksheet,
    min_row,
    max_row,
    min_col,
    max_col
):
    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col
    ):
        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


# ============================================================
# Data Validation
# ============================================================

def add_list_validation(
    worksheet,
    cell_range: str,
    formula: str
):
    validation = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True
    )

    validation.error = (
        "목록에서 값을 선택해주세요."
    )

    validation.errorTitle = (
        "잘못된 값"
    )

    worksheet.add_data_validation(
        validation
    )

    validation.add(
        cell_range
    )


def add_bool_validation(
    worksheet,
    cell_range
):
    add_list_validation(
        worksheet,
        cell_range,
        '"TRUE,FALSE"'
    )


# ============================================================
# README
# ============================================================

def create_readme(workbook):
    ws = workbook.create_sheet(
        "README"
    )

    style_title(
        ws,
        "A1:H1",
        "Story Excel Editor"
    )

    rows = [
        (
            "목적",
            "스토리별 시트를 실제 스토리 작성 도구로 사용합니다."
        ),
        (
            "데이터 흐름",
            "StoryData.asset → CSV → 이 Excel → CSV → StoryData.asset"
        ),
        (
            "Nodes",
            "실제 대사, 지문, 선택지 노드를 위에서 아래로 작성합니다."
        ),
        (
            "Choices",
            "NodeType이 Choice인 노드의 선택지를 작성합니다."
        ),
        (
            "Conditions",
            "해당 스토리 자체의 실행 조건입니다."
        ),
        (
            "Effects",
            "노드별 연출 효과입니다."
        ),
        (
            "StoryId",
            "스토리의 영구 식별자입니다. 제목을 수정해도 StoryId는 가급적 변경하지 마세요."
        ),
        (
            "NodeId",
            "스토리 내부 노드의 식별자입니다."
        ),
        (
            "NextNodeId",
            "현재 노드 다음에 이동할 NodeId입니다."
        ),
        (
            "Branch/Scene",
            "Excel 작성 편의를 위한 보조 컬럼입니다. Common/A/B/C 등의 분기명을 적을 수 있습니다."
        ),
        (
            "Import",
            "같은 StoryId의 StoryData가 Unity에 존재하면 새로 만들지 않고 기존 asset 내용을 업데이트합니다."
        ),
    ]

    start_row = 3

    for index, (title, text) in enumerate(
        rows,
        start=start_row
    ):
        ws.cell(
            row=index,
            column=1,
            value=title
        )

        ws.cell(
            row=index,
            column=2,
            value=text
        )

        ws.cell(
            row=index,
            column=1
        ).fill = PatternFill(
            "solid",
            fgColor=COLOR_SECTION
        )

        ws.cell(
            row=index,
            column=1
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=index,
            column=2
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90


# ============================================================
# Enum Sheet
# ============================================================

def create_enum_sheet(workbook):
    ws = workbook.create_sheet(
        "Enums"
    )

    headers = [
        "Category",
        "Value"
    ]

    style_headers(
        ws,
        1,
        headers
    )

    enum_data = []

    for value in STORY_TYPES:
        enum_data.append(
            ("StoryType", value)
        )

    for value in PLAY_TIMINGS:
        enum_data.append(
            ("StoryPlayTiming", value)
        )

    for value in NODE_TYPES:
        enum_data.append(
            ("StoryNodeType", value)
        )

    for value in CONDITION_TYPES:
        enum_data.append(
            ("StoryConditionType", value)
        )

    for value in EFFECT_TYPES:
        enum_data.append(
            ("StoryEffectType", value)
        )

    for value in EFFECT_TARGETS:
        enum_data.append(
            ("StoryEffectTarget", value)
        )

    for row_index, row in enumerate(
        enum_data,
        start=2
    ):
        ws.cell(
            row=row_index,
            column=1,
            value=row[0]
        )

        ws.cell(
            row=row_index,
            column=2,
            value=row[1]
        )

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 32


# ============================================================
# Character Sheet
# ============================================================


def create_character_sheet(
    workbook,
    characters
):
    ws = workbook.create_sheet(
        "Characters"
    )

    headers = [
        "CharacterId",
        "CharacterName",
        "CharacterNameEng",
        "PortraitId",
        "PortraitLabel",
        "SpriteName"
    ]

    style_headers(
        ws,
        1,
        headers
    )

    for row_index, character in enumerate(
        characters,
        start=2
    ):
        values = [
            get_value(
                character,
                "CharacterId"
            ),
            get_value(
                character,
                "CharacterName"
            ),
            get_value(
                character,
                "CharacterNameEng"
            ),
            get_value(
                character,
                "PortraitId"
            ),
            get_value(
                character,
                "PortraitLabel"
            ),
            get_value(
                character,
                "SpriteName"
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row_index,
                column=column,
                value=value
            )

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["F"].width = 32

    ws.freeze_panes = "A2"

    return len(characters)

# ============================================================
# StoryList
# ============================================================

def create_story_list(
    workbook,
    stories,
    sheet_map
):
    ws = workbook.create_sheet(
        "StoryList"
    )

    headers = [
        "SheetName",
        "StoryId",
        "StoryType",
        "StoryTitle",
        "Year",
        "Month",
        "Day",
        "Priority",
        "PlayOnce",
        "PlayTiming",
        "StartNodeId",
        "Memo",
    ]

    style_headers(
        ws,
        1,
        headers
    )

    for row_index, story in enumerate(
        stories,
        start=2
    ):
        story_id = get_value(
            story,
            "StoryId"
        )

        values = [
            sheet_map.get(
                story_id,
                ""
            ),
            story_id,
            get_value(
                story,
                "StoryType"
            ),
            get_value(
                story,
                "StoryTitle"
            ),
            parse_int(
                get_value(
                    story,
                    "Year"
                )
            ),
            parse_int(
                get_value(
                    story,
                    "Month"
                )
            ),
            parse_int(
                get_value(
                    story,
                    "Day"
                )
            ),
            parse_int(
                get_value(
                    story,
                    "Priority"
                ),
                100
            ),
            normalize_bool(
                get_value(
                    story,
                    "PlayOnce"
                ),
                "TRUE"
            ),
            get_value(
                story,
                "PlayTiming"
            ),
            get_value(
                story,
                "StartNodeId"
            ),
            get_value(
                story,
                "Memo"
            ),
        ]

        for col, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row_index,
                column=col,
                value=value
            )

    add_list_validation(
        ws,
        "C2:C500",
        '"Normal,Force,Unlock,Branch,Event"'
    )

    add_bool_validation(
        ws,
        "I2:I500"
    )

    add_list_validation(
        ws,
        "J2:J500",
        '"BeforeMonthly,Monthly,AfterMonthly"'
    )

    widths = [
        28,
        22,
        16,
        32,
        10,
        10,
        10,
        12,
        12,
        20,
        24,
        36,
    ]

    for index, width in enumerate(
        widths,
        start=1
    ):
        ws.column_dimensions[
            chr(64 + index)
        ].width = width

    ws.freeze_panes = "A2"


# ============================================================
# Story Sheet
# ============================================================

NODE_START_ROW = 12
NODE_MAX_ROW = 220

CHOICE_TITLE_ROW = 224
CHOICE_HEADER_ROW = 225
CHOICE_START_ROW = 226
CHOICE_MAX_ROW = 300

CONDITION_TITLE_ROW = 304
CONDITION_HEADER_ROW = 305
CONDITION_START_ROW = 306
CONDITION_MAX_ROW = 340

EFFECT_TITLE_ROW = 344
EFFECT_HEADER_ROW = 345
EFFECT_START_ROW = 346
EFFECT_MAX_ROW = 400


def create_story_sheet(
    workbook,
    sheet_name,
    story,
    nodes,
    choices,
    conditions,
    effects,
    character_count
):
    ws = workbook.create_sheet(
        sheet_name
    )

    style_title(
        ws,
        "A1:N1",
        f"Story Editor — {get_value(story, 'StoryTitle')}"
    )

    # --------------------------------------------------------
    # 기본 정보
    # --------------------------------------------------------

    style_section(
        ws,
        3,
        "스토리 기본 정보"
    )

    meta_headers = [
        "StoryId",
        "StoryType",
        "StoryTitle",
        "Year",
        "Month",
        "Day",
        "Priority",
        "PlayOnce",
        "PlayTiming",
        "UnlockFactionId",
        "StartNodeId",
        "Memo",
    ]

    style_headers(
        ws,
        4,
        meta_headers
    )

    meta_values = [
        get_value(story, "StoryId"),
        get_value(story, "StoryType"),
        get_value(story, "StoryTitle"),
        parse_int(
            get_value(
                story,
                "Year"
            )
        ),
        parse_int(
            get_value(
                story,
                "Month"
            )
        ),
        parse_int(
            get_value(
                story,
                "Day"
            )
        ),
        parse_int(
            get_value(
                story,
                "Priority"
            ),
            100
        ),
        normalize_bool(
            get_value(
                story,
                "PlayOnce"
            ),
            "TRUE"
        ),
        get_value(
            story,
            "PlayTiming"
        ),
        get_value(
            story,
            "UnlockFactionId"
        ),
        get_value(
            story,
            "StartNodeId"
        ),
        get_value(
            story,
            "Memo"
        ),
    ]

    for column, value in enumerate(
        meta_values,
        start=1
    ):
        ws.cell(
            row=5,
            column=column,
            value=value
        )

    add_list_validation(
        ws,
        "B5",
        '"Normal,Force,Unlock,Branch,Event"'
    )

    add_bool_validation(
        ws,
        "H5"
    )

    add_list_validation(
        ws,
        "I5",
        '"BeforeMonthly,Monthly,AfterMonthly"'
    )

    # 설명
    ws.merge_cells(
        "A7:N8"
    )

    ws["A7"] = (
        "NodeType / Character / Portrait / TRUE·FALSE 값은 "
        "드롭다운으로 선택할 수 있습니다. "
        "Branch/Scene은 Common, A, B, C 등의 분기를 "
        "보기 쉽게 구분하기 위한 Excel 전용 보조 컬럼입니다."
    )

    ws["A7"].fill = PatternFill(
        "solid",
        fgColor=COLOR_INFO
    )

    ws["A7"].alignment = Alignment(
        wrap_text=True,
        vertical="center"
    )

    # --------------------------------------------------------
    # Nodes
    # --------------------------------------------------------

    style_section(
        ws,
        10,
        "Nodes — 실제 대사 / 지문 / 선택지 노드"
    )

    node_headers = [
        "Order",
        "NodeId",
        "NodeType",
        "Text",
        "CharacterId",
        "PortraitId",
        "NextNodeId",
        "DimPortrait",
        "KeepPortrait",
        "UseTypingEffect",
        "AutoAdvance",
        "AutoAdvanceDelay",
        "Branch/Scene",
        "Memo",
    ]

    style_headers(
        ws,
        11,
        node_headers
    )

    nodes = sorted(
        nodes,
        key=lambda row:
            parse_int(
                get_value(
                    row,
                    "NodeOrder"
                )
            )
    )

    for offset, node in enumerate(
        nodes
    ):
        row = NODE_START_ROW + offset

        if row > NODE_MAX_ROW:
            raise ValueError(
                f"{sheet_name}: Node가 "
                f"{NODE_MAX_ROW - NODE_START_ROW + 1}개를 "
                "초과했습니다."
            )

        values = [
            parse_int(
                get_value(
                    node,
                    "NodeOrder"
                )
            ),
            get_value(
                node,
                "NodeId"
            ),
            get_value(
                node,
                "NodeType"
            ),
            get_value(
                node,
                "Text"
            ),
            get_value(
                node,
                "CharacterId"
            ),
            get_value(
                node,
                "PortraitId"
            ),
            get_value(
                node,
                "NextNodeId"
            ),
            normalize_bool(
                get_value(
                    node,
                    "DimPortrait"
                )
            ),
            normalize_bool(
                get_value(
                    node,
                    "KeepPortrait"
                ),
                "TRUE"
            ),
            normalize_bool(
                get_value(
                    node,
                    "UseTypingEffect"
                ),
                "TRUE"
            ),
            normalize_bool(
                get_value(
                    node,
                    "AutoAdvance"
                )
            ),
            parse_float(
                get_value(
                    node,
                    "AutoAdvanceDelay"
                ),
                1.5
            ),
            get_value(
                node,
                "Branch/Scene"
            ),
            get_value(
                node,
                "Memo"
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row,
                column=column,
                value=value
            )

    # 드롭다운
    add_list_validation(
        ws,
        f"C{NODE_START_ROW}:C{NODE_MAX_ROW}",
        '"CharacterDialogue,PlayerDialogue,Narration,Choice,End"'
    )

    for column in (
        "H",
        "I",
        "J",
        "K",
    ):
        add_bool_validation(
            ws,
            f"{column}{NODE_START_ROW}:{column}{NODE_MAX_ROW}"
        )

    if character_count > 0:

        max_char_row = (
            character_count + 1
        )

        add_list_validation(
            ws,
            f"E{NODE_START_ROW}:E{NODE_MAX_ROW}",
            f"'Characters'!$A$2:$A${max_char_row}"
        )

        add_list_validation(
            ws,
            f"F{NODE_START_ROW}:F{NODE_MAX_ROW}",
            f"'Characters'!$C$2:$C${max_char_row}"
        )

    # --------------------------------------------------------
    # 노드 타입별 색상
    # --------------------------------------------------------

    type_rules = [
        (
            "CharacterDialogue",
            COLOR_CHARACTER
        ),
        (
            "PlayerDialogue",
            COLOR_PLAYER
        ),
        (
            "Narration",
            COLOR_NARRATION
        ),
        (
            "Choice",
            COLOR_CHOICE
        ),
        (
            "End",
            COLOR_END
        ),
    ]

    for node_type, color in type_rules:

        rule = FormulaRule(
            formula=[
                f'$C{NODE_START_ROW}="{node_type}"'
            ],
            fill=PatternFill(
                "solid",
                fgColor=color
            )
        )

        ws.conditional_formatting.add(
            f"A{NODE_START_ROW}:N{NODE_MAX_ROW}",
            rule
        )

    # CharacterDialogue인데 CharacterId 없음
    warning_character = FormulaRule(
        formula=[
            (
                f'AND('
                f'$C{NODE_START_ROW}="CharacterDialogue",'
                f'$E{NODE_START_ROW}=""'
                f')'
            )
        ],
        fill=PatternFill(
            "solid",
            fgColor=COLOR_WARNING
        )
    )

    ws.conditional_formatting.add(
        f"E{NODE_START_ROW}:E{NODE_MAX_ROW}",
        warning_character
    )

    # --------------------------------------------------------
    # Choices
    # --------------------------------------------------------

    style_section(
        ws,
        CHOICE_TITLE_ROW,
        "Choices — Choice 노드의 선택지"
    )

    choice_headers = [
        "ChoiceNodeId",
        "ChoiceIndex",
        "ChoiceText",
        "TargetNodeId",
        "ResultKey",
        "ResultValue",
        "UseCondition",
        "RequiredKey",
        "RequiredValue",
        "HideWhenLocked",
        "LockedText",
        "Memo",
    ]

    style_headers(
        ws,
        CHOICE_HEADER_ROW,
        choice_headers
    )

    choices = sorted(
        choices,
        key=lambda row: (
            get_value(
                row,
                "NodeId"
            ),
            parse_int(
                get_value(
                    row,
                    "ChoiceIndex"
                )
            ),
        )
    )

    for offset, choice in enumerate(
        choices
    ):
        row = (
            CHOICE_START_ROW
            + offset
        )

        if row > CHOICE_MAX_ROW:
            raise ValueError(
                f"{sheet_name}: Choice가 너무 많습니다."
            )

        values = [
            get_value(
                choice,
                "NodeId"
            ),
            parse_int(
                get_value(
                    choice,
                    "ChoiceIndex"
                )
            ),
            get_value(
                choice,
                "ChoiceText"
            ),
            get_value(
                choice,
                "TargetNodeId"
            ),
            get_value(
                choice,
                "ResultKey"
            ),
            get_value(
                choice,
                "ResultValue"
            ),
            normalize_bool(
                get_value(
                    choice,
                    "UseCondition"
                )
            ),
            get_value(
                choice,
                "RequiredKey"
            ),
            get_value(
                choice,
                "RequiredValue"
            ),
            normalize_bool(
                get_value(
                    choice,
                    "HideWhenLocked"
                ),
                "TRUE"
            ),
            get_value(
                choice,
                "LockedText"
            ),
            get_value(
                choice,
                "Memo"
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row,
                column=column,
                value=value
            )

    add_bool_validation(
        ws,
        f"G{CHOICE_START_ROW}:G{CHOICE_MAX_ROW}"
    )

    add_bool_validation(
        ws,
        f"J{CHOICE_START_ROW}:J{CHOICE_MAX_ROW}"
    )

    # --------------------------------------------------------
    # Conditions
    # --------------------------------------------------------

    style_section(
        ws,
        CONDITION_TITLE_ROW,
        "Story Conditions — 이 스토리의 실행 조건"
    )

    condition_headers = [
        "ConditionIndex",
        "ConditionType",
        "Key",
        "Value",
        "IntValue",
        "Memo",
    ]

    style_headers(
        ws,
        CONDITION_HEADER_ROW,
        condition_headers
    )

    conditions = sorted(
        conditions,
        key=lambda row:
            parse_int(
                get_value(
                    row,
                    "ConditionIndex"
                )
            )
    )

    for offset, condition in enumerate(
        conditions
    ):
        row = (
            CONDITION_START_ROW
            + offset
        )

        values = [
            parse_int(
                get_value(
                    condition,
                    "ConditionIndex"
                )
            ),
            get_value(
                condition,
                "ConditionType"
            ),
            get_value(
                condition,
                "Key"
            ),
            get_value(
                condition,
                "Value"
            ),
            parse_int(
                get_value(
                    condition,
                    "IntValue"
                )
            ),
            get_value(
                condition,
                "Memo"
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row,
                column=column,
                value=value
            )

    add_list_validation(
        ws,
        f"B{CONDITION_START_ROW}:B{CONDITION_MAX_ROW}",
        (
            '"None,StoryCompleted,StoryNotCompleted,'
            'FactionIntroduced,FactionNotIntroduced,'
            'RelationshipAtLeast,RelationshipAtMost,'
            'ChoiceEquals"'
        )
    )

    # --------------------------------------------------------
    # Effects
    # --------------------------------------------------------

    style_section(
        ws,
        EFFECT_TITLE_ROW,
        "Effects — 필요할 때만 작성"
    )

    effect_headers = [
        "NodeId",
        "EffectIndex",
        "EffectType",
        "Target",
        "Duration",
        "Strength",
        "WaitForCompletion",
        "Memo",
    ]

    style_headers(
        ws,
        EFFECT_HEADER_ROW,
        effect_headers
    )

    effects = sorted(
        effects,
        key=lambda row: (
            get_value(
                row,
                "NodeId"
            ),
            parse_int(
                get_value(
                    row,
                    "EffectIndex"
                )
            ),
        )
    )

    for offset, effect in enumerate(
        effects
    ):
        row = (
            EFFECT_START_ROW
            + offset
        )

        values = [
            get_value(
                effect,
                "NodeId"
            ),
            parse_int(
                get_value(
                    effect,
                    "EffectIndex"
                )
            ),
            get_value(
                effect,
                "EffectType"
            ),
            get_value(
                effect,
                "Target"
            ),
            parse_float(
                get_value(
                    effect,
                    "Duration"
                )
            ),
            parse_float(
                get_value(
                    effect,
                    "Strength"
                ),
                1.0
            ),
            normalize_bool(
                get_value(
                    effect,
                    "WaitForCompletion"
                )
            ),
            get_value(
                effect,
                "Memo"
            ),
        ]

        for column, value in enumerate(
            values,
            start=1
        ):
            ws.cell(
                row=row,
                column=column,
                value=value
            )

    add_list_validation(
        ws,
        f"C{EFFECT_START_ROW}:C{EFFECT_MAX_ROW}",
        (
            '"None,PortraitFadeIn,PortraitFadeOut,'
            'PortraitShakeSmall,PortraitShakeStrong,'
            'PortraitZoomIn,PortraitZoomOut,'
            'ScreenFlash,ScreenDarken,TextShake"'
        )
    )

    add_list_validation(
        ws,
        f"D{EFFECT_START_ROW}:D{EFFECT_MAX_ROW}",
        (
            '"Portrait,DialoguePanel,DialogueText,'
            'Background,EntireScreen"'
        )
    )

    add_bool_validation(
        ws,
        f"G{EFFECT_START_ROW}:G{EFFECT_MAX_ROW}"
    )

    # --------------------------------------------------------
    # 전체 포맷
    # --------------------------------------------------------

    widths = {
        "A": 11,
        "B": 23,
        "C": 21,
        "D": 68,
        "E": 19,
        "F": 19,
        "G": 23,
        "H": 14,
        "I": 14,
        "J": 18,
        "K": 15,
        "L": 18,
        "M": 18,
        "N": 34,
    }

    for column, width in widths.items():
        ws.column_dimensions[
            column
        ].width = width

    apply_standard_cell_style(
        ws,
        NODE_START_ROW,
        NODE_MAX_ROW,
        1,
        14
    )

    apply_standard_cell_style(
        ws,
        CHOICE_START_ROW,
        CHOICE_MAX_ROW,
        1,
        12
    )

    apply_standard_cell_style(
        ws,
        CONDITION_START_ROW,
        CONDITION_MAX_ROW,
        1,
        6
    )

    apply_standard_cell_style(
        ws,
        EFFECT_START_ROW,
        EFFECT_MAX_ROW,
        1,
        8
    )

    for row in range(
        NODE_START_ROW,
        NODE_MAX_ROW + 1
    ):
        ws.row_dimensions[
            row
        ].height = 32

    ws.freeze_panes = (
        f"A{NODE_START_ROW}"
    )

    ws.auto_filter.ref = (
        f"A11:N{NODE_MAX_ROW}"
    )


# ============================================================
# Template
# ============================================================

def create_template_sheet(
    workbook,
    character_count
):
    fake_story = {
        "StoryId": "New_Story",
        "StoryType": "Normal",
        "StoryTitle": "새 스토리",
        "Year": "0",
        "Month": "0",
        "Day": "0",
        "Priority": "100",
        "PlayOnce": "TRUE",
        "PlayTiming": "Monthly",
        "UnlockFactionId": "",
        "StartNodeId": "NEW_N_001",
        "Memo": "",
    }

    create_story_sheet(
        workbook,
        "STORY_TEMPLATE",
        fake_story,
        [],
        [],
        [],
        [],
        character_count
    )


# ============================================================
# Build
# ============================================================

def build_excel(overwrite=False):

    print("=" * 60)
    print("Story Excel Builder")
    print("=" * 60)

    # --------------------------------------------------------
    # 파일 검사
    # --------------------------------------------------------

    stories = read_csv(
        STORIES_CSV,
        required=True
    )

    nodes = read_csv(
        NODES_CSV,
        required=True
    )

    choices = read_csv(
        CHOICES_CSV,
        required=False
    )

    conditions = read_csv(
        CONDITIONS_CSV,
        required=False
    )

    effects = read_csv(
        EFFECTS_CSV,
        required=False
    )

    characters = read_csv(
        CHARACTERS_CSV,
        required=False
    )

    if not stories:
        raise ValueError(
            "Stories.csv에 스토리가 없습니다."
        )

    # --------------------------------------------------------
    # 기존 Excel 보호
    # --------------------------------------------------------

    EXCEL_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    if EXCEL_PATH.exists():

        if not overwrite:

            print()
            print(
                "[STOP] StoryData.xlsx가 이미 존재합니다."
            )

            print(
                "기존 Excel을 보호하기 위해 "
                "생성을 중단했습니다."
            )

            print()
            print(
                "ScriptableObject/CSV 기준으로 "
                "Excel을 완전히 다시 만들려면:"
            )

            print(
                "python StoryExcelBuilder.py --overwrite"
            )

            return

        print(
            "[WARNING] 기존 StoryData.xlsx를 "
            "덮어씁니다."
        )

    # --------------------------------------------------------
    # StoryId별 데이터 그룹
    # --------------------------------------------------------

    nodes_by_story = defaultdict(
        list
    )

    choices_by_story = defaultdict(
        list
    )

    conditions_by_story = defaultdict(
        list
    )

    effects_by_story = defaultdict(
        list
    )

    for row in nodes:
        nodes_by_story[
            get_value(
                row,
                "StoryId"
            )
        ].append(row)

    for row in choices:
        choices_by_story[
            get_value(
                row,
                "StoryId"
            )
        ].append(row)

    for row in conditions:
        conditions_by_story[
            get_value(
                row,
                "StoryId"
            )
        ].append(row)

    for row in effects:
        effects_by_story[
            get_value(
                row,
                "StoryId"
            )
        ].append(row)

    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    workbook = Workbook()

    # 기본 Sheet 삭제
    default_sheet = workbook.active
    workbook.remove(
        default_sheet
    )

    create_readme(
        workbook
    )

    create_enum_sheet(
        workbook
    )

    character_count = (
        create_character_sheet(
            workbook,
            characters
        )
    )

    # --------------------------------------------------------
    # 시트 이름 먼저 결정
    # --------------------------------------------------------

    sheet_map = {}

    used_names = set(
        workbook.sheetnames
    )

    for story in stories:

        story_id = get_value(
            story,
            "StoryId"
        )

        desired = auto_story_sheet_name(
            story
        )

        candidate = desired
        number = 2

        while candidate in used_names:

            suffix = f"_{number}"

            candidate = (
                desired[
                    :31 - len(suffix)
                ]
                + suffix
            )

            number += 1

        used_names.add(
            candidate
        )

        sheet_map[
            story_id
        ] = candidate

    create_story_list(
        workbook,
        stories,
        sheet_map
    )

    create_template_sheet(
        workbook,
        character_count
    )

    # --------------------------------------------------------
    # 실제 Story Sheet
    # --------------------------------------------------------

    for story in stories:

        story_id = get_value(
            story,
            "StoryId"
        )

        sheet_name = sheet_map[
            story_id
        ]

        create_story_sheet(
            workbook,
            sheet_name,
            story,
            nodes_by_story[
                story_id
            ],
            choices_by_story[
                story_id
            ],
            conditions_by_story[
                story_id
            ],
            effects_by_story[
                story_id
            ],
            character_count
        )

        print(
            f"[Story] {story_id}"
            f" → {sheet_name}"
        )

    # --------------------------------------------------------
    # 저장
    # --------------------------------------------------------

    workbook.save(
        EXCEL_PATH
    )

    print()
    print("=" * 60)
    print("Excel 생성 완료")
    print("=" * 60)

    print(
        f"Stories   : {len(stories)}"
    )

    print(
        f"Nodes     : {len(nodes)}"
    )

    print(
        f"Choices   : {len(choices)}"
    )

    print(
        f"Conditions: {len(conditions)}"
    )

    print(
        f"Effects   : {len(effects)}"
    )

    print()

    print(
        f"Output:\n{EXCEL_PATH}"
    )


# ============================================================
# Main
# ============================================================

def main():

    parser = argparse.ArgumentParser(
        description=(
            "Story CSV를 스토리별 Excel 시트로 변환합니다."
        )
    )

    parser.add_argument(
        "--overwrite",
        action="store_true",
        help=(
            "기존 StoryData.xlsx를 덮어씁니다."
        )
    )

    args = parser.parse_args()

    try:

        build_excel(
            overwrite=args.overwrite
        )

    except Exception as exception:

        print()
        print(
            "[ERROR] Excel 생성 실패"
        )

        print(
            str(exception)
        )

        sys.exit(1)


if __name__ == "__main__":
    main()


# 실행 : python StoryExcelBuilder.py